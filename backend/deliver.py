"""Real deliverables — Excel, Word, PowerPoint, Markdown — rendered from a declarative spec.

A local 8B model gets a multi-step stateful API (create -> add_heading -> add_row ->
save) wrong often enough to leave corrupt half-written files behind, but it emits one
JSON blob reliably. So the model only ever produces a spec; this module validates it and
renders it deterministically. Every validation failure is a sentence the model can read
and act on, never a traceback, and nothing touches the target path until the spec is
known good — a bad spec must never produce a corrupt file.
"""
import os
import re
import tempfile
from datetime import date, datetime
from pathlib import Path

import xlsxwriter
from docx import Document
from pptx import Presentation
from pptx.util import Inches as PInches, Pt as PPt
from xlsxwriter.utility import xl_col_to_name

NUMFMT = {
    "text": None,
    "number": "#,##0.00",
    "integer": "#,##0",
    "currency": "$#,##0.00",
    "percent": "0.0%",
    "date": "yyyy-mm-dd",
}
NUMERIC = ("number", "integer", "currency", "percent")
# JSON-Schema-ish type words a model writes when it thinks in terms of data types rather
# than cell formatting. Anything unrecognised falls back to text, which is never wrong.
_TYPE_ALIASES = {
    "string": "text", "str": "text", "text": "text",
    "number": "number", "float": "number", "double": "number", "decimal": "number",
    "int": "integer", "integer": "integer",
    "currency": "currency", "money": "currency",
    "percent": "percent", "percentage": "percent",
    "date": "date", "datetime": "date",
}
BLOCKS = ("heading", "paragraph", "bullets", "numbered", "table", "pagebreak", "quote")
LAYOUTS = ("title", "bullets", "table", "text")
MAX_COL_WIDTH = 60
# python-pptx has no autofit or reflow: text past the placeholder simply runs off the
# slide and nobody notices until the meeting. These are the volumes where 18pt stops fitting.
MAX_BULLETS = 9
MAX_BULLET_CHARS = 180


class SpecError(ValueError):
    """Message is guidance aimed at the model that produced the spec, not a stack trace."""


# ---------------------------------------------------------------- validation helpers

def _obj(spec, kind: str) -> dict:
    if not isinstance(spec, dict):
        raise SpecError(f"{kind} spec: expected a JSON object, got {type(spec).__name__}.")
    return spec


def _listify(v) -> list:
    """A model asked for a list routinely emits the bare item; wrapping it is never ambiguous."""
    if v is None:
        return []
    return list(v) if isinstance(v, (list, tuple)) else [v]


def _cell(v):
    """Unwrap a cell the model described instead of stated.

    Models routinely emit {"value": "BOLT-M12", "formula": false} or {"text": "..."} for a
    cell. Left alone these get stringified into the sheet, producing a file that reports
    the right row count and contains "{'formula': False, 'value': 'BOLT-M12'}" — a success
    that is actually corrupt, which is the worst possible outcome. The wrapper is
    unambiguous, so unwrap it.
    """
    if isinstance(v, dict):
        for key in ("value", "text", "content", "v"):
            if key in v:
                return v[key]
    return v


def _rows(raw, ncols: int, where: str, headers: list[str] | None = None) -> list[list]:
    if raw is None:
        return []
    if not isinstance(raw, (list, tuple)):
        raise SpecError(f"{where}: 'rows' must be a list of rows, e.g. [[\"EMEA\", 1240000]], got {type(raw).__name__}.")
    raw = list(raw)
    # A single flat row where a list-of-rows was expected is only unambiguous when it
    # matches the column count exactly; anything else is the model's to fix.
    if ncols > 1 and raw and not any(isinstance(r, (list, tuple)) for r in raw) and len(raw) == ncols:
        raw = [raw]
    out = []
    for i, row in enumerate(raw):
        # {"Part": "BOLT-M12", "Qty": 40} is the other shape models reach for, and the
        # headers tell us the column order unambiguously, so honour it.
        if isinstance(row, dict) and headers and not any(k in row for k in ("value", "text", "v")):
            lookup = {str(k).strip().lower(): v for k, v in row.items()}
            row = [lookup.get(h.strip().lower()) for h in headers]
        row = list(row) if isinstance(row, (list, tuple)) else [row]
        row = [_cell(v) for v in row]
        if len(row) != ncols:
            raise SpecError(
                f"{where} has {ncols} columns but row {i} has {len(row)} values. "
                "Every row must match the column count."
            )
        out.append(row)
    return out


def _flatten(items, where: str) -> list[tuple[int, str]]:
    """["a", ["sub"], "b"] -> [(0,"a"), (1,"sub"), (0,"b")] — one nesting level, which is
    all any of the three formats renders anyway."""
    flat = []
    for item in _listify(items):
        if isinstance(item, (list, tuple)):
            flat += [(1, str(x)) for x in item if x is not None]
        elif item is not None:
            flat.append((0, str(item)))
    if not flat:
        raise SpecError(f"{where} has no items. Expected a list of strings, e.g. [\"first\", \"second\"].")
    return flat


_BAD_SHEET = re.compile(r"[\[\]:*?/\\]")


def _sheet_name(name, used: set) -> tuple[str, str]:
    clean = _BAD_SHEET.sub("-", str(name)).strip("'")[:31].strip() or "Sheet"
    base, n = clean, 2
    while clean.lower() in used:
        tail = f" ({n})"
        clean, n = base[:31 - len(tail)] + tail, n + 1
    used.add(clean.lower())
    return clean, (f"renamed sheet {str(name)!r} to {clean!r} for Excel" if clean != str(name) else "")


def _norm_xlsx(spec) -> tuple[list[dict], list[str]]:
    _obj(spec, "xlsx")
    # A model asked for a one-sheet workbook writes {"columns":…, "rows":…} far more
    # often than the nested form, and rejecting that costs a whole extra turn on a model
    # that may not recover. Wrapping it is unambiguous, so meet it where it is.
    if "sheets" not in spec and ("columns" in spec or "rows" in spec):
        spec = {"sheets": [{k: v for k, v in spec.items() if k != "sheets"}]}
    if "sheets" not in spec or spec["sheets"] is None:
        raise SpecError(
            'xlsx spec: missing required key \'sheets\'. Expected '
            '{"sheets": [{"name": str, "columns": [...], "rows": [[...]]}]}'
        )
    sheets = _listify(spec["sheets"])
    if not sheets:
        raise SpecError("xlsx spec: 'sheets' is empty. Provide at least one sheet with 'columns' and 'rows'.")

    used, out, notes = set(), [], []
    for i, sheet in enumerate(sheets):
        if not isinstance(sheet, dict):
            raise SpecError(f"xlsx spec: sheet {i} must be an object with 'name', 'columns' and 'rows', got {type(sheet).__name__}.")
        cols = _listify(sheet.get("columns"))
        if not cols:
            raise SpecError(f'xlsx spec: sheet {i} has no \'columns\'. Expected [{{"header": "Region", "format": "text"}}, ...].')
        norm = []
        for j, col in enumerate(cols):
            if isinstance(col, str):
                col = {"header": col}
            # "name"/"title" for the header and "type" for the format are the spellings a
            # model reaches for unprompted; both map cleanly onto ours.
            if isinstance(col, dict) and not col.get("header"):
                col = {**col, "header": col.get("name") or col.get("title") or ""}
            if isinstance(col, dict) and not col.get("format") and col.get("type"):
                col = {**col, "format": _TYPE_ALIASES.get(str(col["type"]).lower(), "text")}
            if not isinstance(col, dict) or not col.get("header"):
                raise SpecError(f'xlsx spec: sheet {i} column {j} must be a header string or {{"header": ...}}, got {col!r}.')
            fmt = str(col.get("format") or "text").lower()
            if fmt not in NUMFMT:
                raise SpecError(f"xlsx spec: sheet {i} column {j} has unknown format {fmt!r}. Valid formats: {', '.join(NUMFMT)}.")
            norm.append({"header": str(col["header"]), "format": fmt, "width": col.get("width"),
                         "explicit": bool(col.get("format") or col.get("type"))})
        name, note = _sheet_name(sheet.get("name") or f"Sheet{i + 1}", used)
        if note:
            notes.append(note)
        data = _rows(sheet.get("rows"), len(norm), f"xlsx spec: sheet {i}",
                     [c["header"] for c in norm])
        # A sheet with headers and no data is almost never what was asked for; it means
        # the model described the columns and dropped the values. Writing it anyway
        # produces an empty file that the model then confidently tells the user is full,
        # which is far worse than one more turn. Refuse, and say what's missing.
        if not data:
            raise SpecError(
                f"xlsx spec: sheet {i} ({name!r}) has columns but no rows, so the file would "
                f"be empty. Put the actual data in 'rows' as a list of lists with "
                f"{len(norm)} values each, e.g. "
                f"{[[c['header'] + '_value' for c in norm]]}."
            )
        # A column the model didn't type, whose every value is a number, is a number
        # column. Left as text it lands left-aligned and Excel won't SUM it — a
        # spreadsheet you can't do arithmetic in isn't really a spreadsheet.
        for c, col in enumerate(norm):
            if col["explicit"]:
                continue
            vals = [r[c] for r in data if r[c] not in (None, "")]
            # A leading zero means it's an identifier, not a quantity — part numbers, zip
            # codes and account numbers all lose their meaning the moment Excel treats
            # them as numbers, so those stay text.
            if any(isinstance(v, str) and len(v) > 1 and v.strip().startswith("0") for v in vals):
                continue
            nums = [_num(v, "number") for v in vals]
            if vals and all(n is not None for n in nums):
                col["format"] = "integer" if all(float(n).is_integer() for n in nums) else "number"

        out.append({
            "name": name,
            "columns": norm,
            "rows": data,
            "total_row": bool(sheet.get("total_row")),
            "freeze_header": sheet.get("freeze_header", True),
        })
    return out, notes


def _norm_blocks(spec, kind: str) -> list[dict]:
    _obj(spec, kind)
    if "blocks" not in spec or spec["blocks"] is None:
        raise SpecError(
            f'{kind} spec: missing required key \'blocks\'. Expected '
            f'{{"title": str, "blocks": [{{"type": "heading", "level": 2, "text": "..."}}, ...]}}. '
            f"Valid types: {', '.join(BLOCKS)}."
        )
    out = []
    for i, block in enumerate(_listify(spec["blocks"])):
        if isinstance(block, str):
            block = {"type": "paragraph", "text": block}
        if not isinstance(block, dict):
            raise SpecError(f"{kind} spec: block {i} must be an object like {{\"type\": \"paragraph\", \"text\": \"...\"}}, got {type(block).__name__}.")
        btype = str(block.get("type") or "").lower().strip()
        if btype not in BLOCKS:
            raise SpecError(f"{kind} spec: block {i} has unknown type {btype!r}. Valid types: {', '.join(BLOCKS)}.")
        norm = {"type": btype}
        if btype in ("heading", "paragraph", "quote"):
            if not str(block.get("text") or "").strip():
                raise SpecError(f"{kind} spec: block {i} ({btype}) is missing 'text'.")
            norm["text"] = str(block["text"])
            norm["level"] = max(1, min(4, int(block.get("level", 2) or 2))) if btype == "heading" else 0
        elif btype in ("bullets", "numbered"):
            norm["items"] = _flatten(block.get("items"), f"{kind} spec: block {i} ({btype})")
        elif btype == "table":
            headers = [str(h) for h in _listify(block.get("headers"))]
            if not headers:
                raise SpecError(f'{kind} spec: block {i} (table) needs \'headers\', e.g. {{"headers": ["Region", "Rev"], "rows": [["EMEA", "1.2M"]]}}.')
            norm["headers"] = headers
            norm["rows"] = _rows(block.get("rows"), len(headers), f"{kind} spec: block {i} table", headers)
        out.append(norm)
    if not out:
        raise SpecError(f"{kind} spec: 'blocks' is empty. Provide at least one block, e.g. {{\"type\": \"paragraph\", \"text\": \"...\"}}.")
    return out


def _norm_slides(spec) -> list[dict]:
    _obj(spec, "pptx")
    if "slides" not in spec or spec["slides"] is None:
        raise SpecError(
            'pptx spec: missing required key \'slides\'. Expected '
            '{"slides": [{"layout": "bullets", "title": "...", "bullets": [...]}]}. '
            f"Valid layouts: {', '.join(LAYOUTS)}."
        )
    slides = _listify(spec["slides"])
    if not slides:
        raise SpecError("pptx spec: 'slides' is empty. Provide at least one slide.")
    out = []
    for i, slide in enumerate(slides):
        if not isinstance(slide, dict):
            raise SpecError(f"pptx spec: slide {i} must be an object with 'layout' and 'title', got {type(slide).__name__}.")
        layout = str(slide.get("layout") or "").lower().strip()
        if not layout:  # inferring from the keys present beats rejecting a usable slide
            layout = "bullets" if slide.get("bullets") else "table" if slide.get("headers") else "text" if slide.get("text") else "title"
        if layout not in LAYOUTS:
            raise SpecError(f"pptx spec: slide {i} has unknown layout {layout!r}. Valid layouts: {', '.join(LAYOUTS)}.")
        norm = {"layout": layout, "title": str(slide.get("title") or ""), "notes": slide.get("notes")}
        if layout == "bullets":
            norm["bullets"] = _flatten(slide.get("bullets"), f"pptx spec: slide {i} (bullets)")
        elif layout == "table":
            headers = [str(h) for h in _listify(slide.get("headers"))]
            if not headers:
                raise SpecError(f'pptx spec: slide {i} (table) needs \'headers\', e.g. {{"headers": ["Region", "Rev"], "rows": [["EMEA", "1.2M"]]}}.')
            norm["headers"] = headers
            norm["rows"] = _rows(slide.get("rows"), len(headers), f"pptx spec: slide {i} table", headers)
        else:
            norm["text"] = str(slide.get("text") or slide.get("subtitle") or "")
        out.append(norm)
    return out


# ---------------------------------------------------------------- writing

def _atomic(path: str, render) -> str:
    """Render into a temp file beside the target, then os.replace() it in: a crash — or
    the model retrying the same call — can never leave a half-written .xlsx at `path`."""
    path = os.path.abspath(path)
    parent = os.path.dirname(path) or "."
    os.makedirs(parent, exist_ok=True)
    fd, tmp = tempfile.mkstemp(dir=parent, prefix=".ada-", suffix=os.path.splitext(path)[1])
    os.close(fd)
    try:
        render(tmp)
        os.replace(tmp, path)
    except Exception:
        try:
            os.unlink(tmp)
        except OSError:
            pass
        raise
    return path


def _done(path: str, detail: str, notes=()) -> str:
    summary = f"Wrote {path} — {detail}, {os.path.getsize(path) / 1024:.1f} KB"
    return summary + (" | " + "; ".join(notes) if notes else "")


def _num(v, fmt: str):
    """'1,240,000', '$1.2', '(500)' and '31%' are all things a model puts in a numeric
    column; converting them is safe. Anything else stays text rather than becoming a lie."""
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return v
    s = str(v).strip().replace(",", "").replace("$", "").replace("€", "").replace("£", "")
    pct = s.endswith("%")
    s = s[:-1] if pct else s
    if s.startswith("(") and s.endswith(")"):
        s = "-" + s[1:-1]
    try:
        n = float(s)
    except ValueError:
        return None
    return n / 100 if pct and fmt == "percent" else n


def _date(v):
    if isinstance(v, (datetime, date)):
        return v
    try:
        return datetime.fromisoformat(str(v).strip())
    except ValueError:
        return None



_MD_TABLE_SEP = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")


def blocks_from_markdown(text: str) -> list[dict]:
    """Markdown -> the block vocabulary the docx/md/pptx writers already speak.

    Deliberately small: headings, paragraphs, bullets, numbered lists, block quotes and
    pipe tables are the whole of what a model writes for a letter, memo or report.
    Anything unrecognised stays a paragraph, which is always a safe rendering.
    """
    blocks: list[dict] = []
    lines = str(text).replace("\r\n", "\n").split("\n")
    i = 0
    para: list[str] = []

    def flush() -> None:
        if para:
            blocks.append({"type": "paragraph", "text": "\n".join(para).strip()})
            para.clear()

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            flush()
            i += 1
            continue

        heading = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if heading:
            flush()
            blocks.append({"type": "heading", "level": min(len(heading.group(1)), 3),
                           "text": heading.group(2).strip()})
            i += 1
            continue

        if stripped in ("---", "***", "___"):
            flush()
            i += 1  # a horizontal rule has no counterpart in the block vocabulary
            continue

        # A pipe table needs its separator row to be a table rather than prose.
        if stripped.startswith("|") and i + 1 < len(lines) and _MD_TABLE_SEP.match(lines[i + 1]):
            flush()
            cells = lambda r: [c.strip() for c in r.strip().strip("|").split("|")]
            headers = cells(stripped)
            rows, i = [], i + 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                row = cells(lines[i])
                # Pad or trim rather than reject: a ragged row in otherwise good output
                # should not lose the whole table.
                rows.append((row + [""] * len(headers))[: len(headers)])
                i += 1
            blocks.append({"type": "table", "headers": headers, "rows": rows})
            continue

        if re.match(r"^[-*+]\s+", stripped):
            flush()
            items = []
            while i < len(lines) and re.match(r"^\s*[-*+]\s+", lines[i]):
                items.append(re.sub(r"^\s*[-*+]\s+", "", lines[i]).strip())
                i += 1
            blocks.append({"type": "bullets", "items": items})
            continue

        if re.match(r"^\d+[.)]\s+", stripped):
            flush()
            items = []
            while i < len(lines) and re.match(r"^\s*\d+[.)]\s+", lines[i]):
                items.append(re.sub(r"^\s*\d+[.)]\s+", "", lines[i]).strip())
                i += 1
            blocks.append({"type": "numbered", "items": items})
            continue

        if stripped.startswith(">"):
            flush()
            quote = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote.append(lines[i].strip().lstrip(">").strip())
                i += 1
            blocks.append({"type": "quote", "text": " ".join(quote)})
            continue

        para.append(stripped)
        i += 1

    flush()
    return blocks


def spec_from_markdown(text: str, kind: str) -> dict:
    """Wrap Markdown into the spec shape for `kind`. The first heading becomes the title
    when there is one, so a letter does not end up with a duplicated header."""
    blocks = blocks_from_markdown(text)
    title = ""
    if blocks and blocks[0]["type"] == "heading":
        title = blocks[0]["text"]
        blocks = blocks[1:]
    if not blocks:
        blocks = [{"type": "paragraph", "text": str(text).strip() or "(empty)"}]

    if kind == "pptx":
        # One slide per heading; loose prose before the first heading becomes the opener.
        slides, current = [], None
        for b in blocks:
            if b["type"] == "heading":
                if current:
                    slides.append(current)
                current = {"layout": "bullets", "title": b["text"], "bullets": []}
            elif current is None:
                slides.append({"layout": "text", "title": title or "Overview",
                               "text": b.get("text", "")})
            elif b["type"] in ("bullets", "numbered"):
                current["bullets"].extend(b["items"])
            elif b["type"] == "table":
                if current["bullets"]:
                    slides.append(current)
                    current = None
                slides.append({"layout": "table", "title": b.get("headers", [""])[0] or "Table",
                               "headers": b["headers"], "rows": b["rows"]})
            elif b.get("text"):
                current["bullets"].append(b["text"])
        if current:
            slides.append(current)
        return {"title": title or "Presentation", "slides": slides or
                [{"layout": "text", "title": title or "Notes", "text": str(text)[:500]}]}

    return {"title": title, "blocks": blocks}

def write_xlsx(path: str, spec: dict) -> str:
    sheets, notes = _norm_xlsx(spec)
    total_rows = sum(len(s["rows"]) for s in sheets)

    def render(tmp):
        wb = xlsxwriter.Workbook(tmp, {"constant_memory": False})
        header_fmt = wb.add_format({"bold": True, "bg_color": "#DDEBF7", "bottom": 1})
        cell_fmts = {k: wb.add_format({"num_format": v} if v else {}) for k, v in NUMFMT.items()}
        total_fmts = {k: wb.add_format({"bold": True, "top": 1, **({"num_format": v} if v else {})}) for k, v in NUMFMT.items()}

        for sheet in sheets:
            ws = wb.add_worksheet(sheet["name"])
            cols, rows = sheet["columns"], sheet["rows"]
            widths = [len(c["header"]) + 4 for c in cols]
            sums = [0.0] * len(cols)

            for c, col in enumerate(cols):
                ws.write_string(0, c, col["header"], header_fmt)

            for r, row in enumerate(rows, start=1):
                for c, val in enumerate(row):
                    fmt = cols[c]["format"]
                    widths[c] = max(widths[c], len(str("" if val is None else val)) + 2)
                    if val is None or val == "":
                        ws.write_blank(r, c, None, cell_fmts[fmt])
                    elif fmt in NUMERIC:
                        n = _num(val, fmt)
                        if n is None:
                            ws.write_string(r, c, str(val))
                        else:
                            ws.write_number(r, c, n, cell_fmts[fmt])
                            sums[c] += n
                    elif fmt == "date" and _date(val) is not None:
                        ws.write_datetime(r, c, _date(val), cell_fmts[fmt])
                    else:
                        ws.write_string(r, c, str(val), cell_fmts[fmt])

            last = len(rows)
            if sheet["total_row"] and rows:
                t = last + 1
                ws.write_string(t, 0, "Total", total_fmts["text"])
                for c, col in enumerate(cols):
                    if c == 0 or col["format"] not in NUMERIC:
                        continue
                    letter = xl_col_to_name(c)
                    # Cache the computed value alongside the formula: Excel recalculates,
                    # but anything reading the file without a calc engine still sees a number.
                    ws.write_formula(t, c, f"=SUM({letter}2:{letter}{last + 1})", total_fmts[col["format"]], sums[c])

            for c, col in enumerate(cols):
                ws.set_column(c, c, col["width"] or min(widths[c], MAX_COL_WIDTH))
            ws.autofilter(0, 0, max(last, 1), len(cols) - 1)
            if sheet["freeze_header"]:
                ws.freeze_panes(1, 0)
        wb.close()

    out = _atomic(path, render)
    return _done(out, f"{len(sheets)} sheets, {total_rows} rows", notes)


def _styled(obj, name: str, fallback: str):
    try:
        obj.style = name
    except KeyError:  # a swapped-in corporate template may not carry the built-in style
        obj.style = fallback


def write_docx(path: str, spec: dict) -> str:
    blocks = _norm_blocks(spec, "docx")
    title, subtitle = str(spec.get("title") or ""), str(spec.get("subtitle") or "")

    def render(tmp):
        doc = Document()
        if title:
            doc.add_heading(title, level=0)
        if subtitle:
            _styled(doc.add_paragraph(subtitle), "Subtitle", "Normal")
        for block in blocks:
            kind = block["type"]
            if kind == "heading":
                doc.add_heading(block["text"], level=block["level"])
            elif kind == "paragraph":
                doc.add_paragraph(block["text"])
            elif kind == "quote":
                _styled(doc.add_paragraph(block["text"]), "Intense Quote", "Normal")
            elif kind in ("bullets", "numbered"):
                base = "List Bullet" if kind == "bullets" else "List Number"
                for level, text in block["items"]:
                    _styled(doc.add_paragraph(text), f"{base} 2" if level else base, "Normal")
            elif kind == "pagebreak":
                doc.add_page_break()
            elif kind == "table":
                table = doc.add_table(rows=1, cols=len(block["headers"]))
                _styled(table, "Light Grid Accent 1", "Table Grid")
                for c, head in enumerate(block["headers"]):
                    cell = table.rows[0].cells[c]
                    cell.text = head
                    cell.paragraphs[0].runs[0].bold = True
                for row in block["rows"]:
                    cells = table.add_row().cells
                    for c, val in enumerate(row):
                        cells[c].text = "" if val is None else str(val)
        doc.save(tmp)

    out = _atomic(path, render)
    words = sum(len(b.get("text", "").split()) for b in blocks)
    return _done(out, f"{len(blocks)} blocks, ~{words + len(title.split())} words")


def _bullet_size(items) -> tuple[int, str]:
    longest = max(len(t) for _, t in items)
    size = 18
    if len(items) > MAX_BULLETS:
        size -= len(items) - MAX_BULLETS
    if longest > MAX_BULLET_CHARS:
        size = min(size, 12)
    size = max(10, size)
    if size == 18:
        return size, ""
    return size, f"{len(items)} bullets / {longest} chars shrunk to {size}pt"


def write_pptx(path: str, spec: dict) -> str:
    slides = _norm_slides(spec)
    deck_title, deck_sub = str(spec.get("title") or ""), str(spec.get("subtitle") or "")
    warnings = []

    def render(tmp):
        prs = Presentation()
        if deck_title and slides[0]["layout"] != "title":
            cover = prs.slides.add_slide(prs.slide_layouts[0])
            cover.shapes.title.text = deck_title
            if deck_sub:
                cover.placeholders[1].text = deck_sub

        for i, s in enumerate(slides):
            layout = {"title": 0, "bullets": 1, "text": 1, "table": 5}[s["layout"]]
            slide = prs.slides.add_slide(prs.slide_layouts[layout])
            if slide.shapes.title is not None:
                slide.shapes.title.text = s["title"] or deck_title

            if s["layout"] == "bullets":
                size, warn = _bullet_size(s["bullets"])
                if warn:
                    warnings.append(f"slide {i} ({s['title'] or 'untitled'}): {warn}")
                body = slide.placeholders[1].text_frame
                body.word_wrap = True
                for n, (level, text) in enumerate(s["bullets"]):
                    para = body.paragraphs[0] if n == 0 else body.add_paragraph()
                    para.text = text
                    para.level = level
                    para.font.size = PPt(size if level == 0 else max(10, size - 2))
            elif s["layout"] == "table":
                nrows = len(s["rows"]) + 1
                shape = slide.shapes.add_table(
                    nrows, len(s["headers"]), PInches(0.5), PInches(1.8),
                    prs.slide_width - PInches(1), PInches(0.4) * nrows,
                )
                table = shape.table
                for c, head in enumerate(s["headers"]):
                    table.cell(0, c).text = head
                    table.cell(0, c).text_frame.paragraphs[0].runs[0].font.bold = True
                for r, row in enumerate(s["rows"], start=1):
                    for c, val in enumerate(row):
                        table.cell(r, c).text = "" if val is None else str(val)
            elif s["text"]:
                slide.placeholders[1].text = s["text"]

            if s["notes"]:
                slide.notes_slide.notes_text_frame.text = str(s["notes"])
        prs.save(tmp)

    out = _atomic(path, render)
    count = len(slides) + (1 if deck_title and slides[0]["layout"] != "title" else 0)
    return _done(out, f"{count} slides", [f"WARNING: {w}, may overflow — split it" for w in warnings])


def write_md(path: str, spec: dict) -> str:
    blocks = _norm_blocks(spec, "md")
    lines = []
    if spec.get("title"):
        lines += [f"# {spec['title']}", ""]
    if spec.get("subtitle"):
        lines += [f"*{spec['subtitle']}*", ""]
    for block in blocks:
        kind = block["type"]
        if kind == "heading":
            lines += [f"{'#' * (block['level'] + 1)} {block['text']}", ""]
        elif kind == "paragraph":
            lines += [block["text"], ""]
        elif kind == "quote":
            lines += [f"> {block['text']}", ""]
        elif kind == "bullets":
            lines += [f"{'  ' * lvl}- {text}" for lvl, text in block["items"]] + [""]
        elif kind == "numbered":
            lines += [f"{'  ' * lvl}1. {text}" for lvl, text in block["items"]] + [""]
        elif kind == "pagebreak":
            lines += ["---", ""]
        elif kind == "table":
            cells = lambda row: "| " + " | ".join(str("" if v is None else v).replace("|", r"\|") for v in row) + " |"
            lines += [cells(block["headers"]), "|" + "---|" * len(block["headers"])]
            lines += [cells(row) for row in block["rows"]] + [""]

    text = "\n".join(lines).rstrip() + "\n"
    # write_text closes the handle before _atomic's os.replace — Windows refuses to
    # replace a file that is still open.
    out = _atomic(path, lambda tmp: Path(tmp).write_text(text, encoding="utf-8"))
    return _done(out, f"{len(blocks)} blocks, {len(lines)} lines")


_WRITERS = {"xlsx": write_xlsx, "docx": write_docx, "pptx": write_pptx, "md": write_md}


def write(kind: str, path: str, spec: dict) -> str:
    writer = _WRITERS.get(str(kind).lower().lstrip("."))
    if writer is None:
        raise SpecError(f"unknown kind {kind!r}. Valid kinds: {', '.join(_WRITERS)}.")
    return writer(path, spec)


# ---------------------------------------------------------------- runnable check

def _raises(fn, *needles) -> str:
    try:
        fn()
    except SpecError as e:
        for needle in needles:
            assert needle in str(e), f"guidance {str(e)!r} is missing {needle!r}"
        return str(e)
    raise AssertionError("expected SpecError, got a written file")


def demo():
    import openpyxl

    out = tempfile.mkdtemp(prefix="ada-deliver-")

    # ---- xlsx
    xlsx = os.path.join(out, "q3.xlsx")
    summary = write("xlsx", xlsx, {"sheets": [
        {"name": "Q3 Summary",
         "columns": [{"header": "Region", "format": "text"},
                     {"header": "Revenue", "width": 14, "format": "currency"},
                     {"header": "Margin", "format": "percent"}],
         "rows": [["EMEA", 1240000, 0.31], ["APAC", "980,000", 0.22], ["AMER", None, 0.18]],
         "total_row": True},
        {"name": "Raw/Data:[2024]*", "columns": ["Item", "Count"], "rows": [["widgets", 3]]},
    ]})
    assert os.path.getsize(xlsx) > 3000, summary
    assert "renamed sheet" in summary and "2 sheets" in summary, summary

    wb = openpyxl.load_workbook(xlsx, data_only=True)
    ws = wb["Q3 Summary"]
    assert ws["A1"].value == "Region" and ws["B2"].value == 1240000
    assert ws["B3"].value == 980000, f"numeric string not coerced: {ws['B3'].value!r}"
    assert ws["A5"].value == "Total", f"total row misplaced: {[c.value for c in ws['A']]}"
    assert ws["B5"].value == 2220000, f"total row wrong: {ws['B5'].value!r}"
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref
    assert "/" not in wb.sheetnames[1] and ":" not in wb.sheetnames[1], wb.sheetnames
    wb.close()

    _raises(lambda: write_xlsx(xlsx, {"title": "nope"}), "missing required key 'sheets'")
    # The shapes a model actually emits for a one-sheet request: no 'sheets' wrapper, and
    # columns described as {"name", "type"} rather than {"header", "format"}.
    flat = os.path.join(out, "flat.xlsx")
    write_xlsx(flat, {"columns": [{"name": "Part", "type": "string"},
                                  {"name": "Qty", "type": "int"}],
                      "rows": [["BOLT-M12", 40], ["BOLT-M8", 25]]})
    fwb = openpyxl.load_workbook(flat)
    fws = fwb.active
    assert [fws["A1"].value, fws["B1"].value] == ["Part", "Qty"], "flat spec headers"
    assert fws["A2"].value == "BOLT-M12" and fws["B2"].value == 40
    fwb.close()

    # Cells described as objects, and rows keyed by header — both observed from a real
    # model run, and both previously stringified into the sheet as Python dict reprs.
    wrapped = os.path.join(out, "wrapped.xlsx")
    write_xlsx(wrapped, {"columns": ["Part", "Qty"], "rows": [
        [{"formula": False, "value": "BOLT-M12"}, {"value": 40}],
        {"Part": "BOLT-M8", "Qty": 25},
    ]})
    wwb = openpyxl.load_workbook(wrapped)
    got = [r for r in wwb.active.iter_rows(min_row=2, values_only=True)]
    wwb.close()
    assert got == [("BOLT-M12", 40), ("BOLT-M8", 25)], f"cell/row unwrapping failed: {got}"

    # Untyped all-numeric column becomes a real number column; an identifier with a
    # leading zero must not, or the part number silently loses it.
    ids = os.path.join(out, "ids.xlsx")
    write_xlsx(ids, {"columns": ["Code", "Qty"], "rows": [["007", 5], ["012", 7]]})
    iwb = openpyxl.load_workbook(ids)
    irows = [r for r in iwb.active.iter_rows(min_row=2, values_only=True)]
    iwb.close()
    assert irows == [("007", 5), ("012", 7)], f"leading zeros lost: {irows}"
    _raises(lambda: write_xlsx(xlsx, {"sheets": [{"columns": ["a", "b", "c"], "rows": [[1, 2, 3], [1, 2, 3, 4, 5]]}]}),
            "3 columns but row 1 has 5 values", "match the column count")
    _raises(lambda: write_xlsx(xlsx, {"sheets": [{"columns": [{"header": "x", "format": "money"}], "rows": []}]}),
            "unknown format 'money'", "currency")
    print("PASS xlsx:", summary)

    # ---- docx
    docx_path = os.path.join(out, "review.docx")
    summary = write("docx", docx_path, {"title": "Q3 Review", "subtitle": "Internal", "blocks": [
        {"type": "heading", "level": 2, "text": "Findings"},
        {"type": "paragraph", "text": "Revenue grew 12% quarter over quarter."},
        {"type": "bullets", "items": ["EMEA led growth", ["Germany +18%"], "APAC flat"]},
        {"type": "numbered", "items": ["Lock the forecast", "Brief the board"]},
        {"type": "table", "headers": ["Region", "Rev"], "rows": [["EMEA", "1.2M"], ["APAC", "0.98M"]]},
        {"type": "pagebreak"},
        {"type": "quote", "text": "We are ahead of plan."},
    ]})
    assert os.path.getsize(docx_path) > 10000, summary
    doc = Document(docx_path)
    styles = {p.style.name: p.text for p in doc.paragraphs}
    assert styles.get("Heading 2") == "Findings", styles
    assert any(p.text.startswith("Revenue grew") for p in doc.paragraphs)
    assert "List Bullet" in styles and "List Number" in styles, styles
    table = doc.tables[0]
    assert table.rows[0].cells[0].text == "Region" and table.rows[1].cells[1].text == "1.2M"
    assert table.rows[0].cells[0].paragraphs[0].runs[0].bold

    _raises(lambda: write_docx(docx_path, {"title": "x"}), "missing required key 'blocks'")
    _raises(lambda: write_docx(docx_path, {"blocks": [{"type": "paragraph", "text": "ok"},
                                                      {"type": "heading", "text": "h"},
                                                      {"type": "bulletlist", "items": ["a"]}]}),
            "block 2 has unknown type 'bulletlist'", "heading, paragraph, bullets, numbered, table, pagebreak, quote")
    print("PASS docx:", summary)

    # ---- pptx
    pptx_path = os.path.join(out, "deck.pptx")
    summary = write("pptx", pptx_path, {"title": "Q3 Review", "slides": [
        {"layout": "title", "title": "Q3 Review", "subtitle": "Internal"},
        {"layout": "bullets", "title": "Findings", "bullets": ["EMEA led", ["Germany +18%"], "APAC flat"],
         "notes": "Keep this to two minutes."},
        {"layout": "table", "title": "Numbers", "headers": ["Region", "Rev"], "rows": [["EMEA", "1.2M"]]},
        {"layout": "text", "title": "Next", "text": "Lock the forecast by Friday."},
    ]})
    assert os.path.getsize(pptx_path) > 20000, summary
    prs = Presentation(pptx_path)
    slides = list(prs.slides)
    assert len(slides) == 4, len(slides)
    assert slides[0].shapes.title.text == "Q3 Review"
    assert slides[1].notes_slide.notes_text_frame.text.startswith("Keep this")
    table = next(sh.table for sh in slides[2].shapes if sh.has_table)
    assert table.cell(0, 0).text == "Region" and table.cell(1, 1).text == "1.2M"
    assert "WARNING" not in summary, summary

    overflow = write_pptx(os.path.join(out, "big.pptx"), {"slides": [
        {"layout": "bullets", "title": "Everything", "bullets": [f"point {i}" for i in range(20)]}]})
    assert "WARNING" in overflow and "Everything" in overflow and "pt" in overflow, overflow

    _raises(lambda: write_pptx(pptx_path, {"title": "x"}), "missing required key 'slides'")
    _raises(lambda: write_pptx(pptx_path, {"slides": [{"layout": "chart", "title": "x"}]}),
            "unknown layout 'chart'", "title, bullets, table, text")
    print("PASS pptx:", summary)
    print("PASS pptx overflow guard:", overflow)

    # ---- md
    md_path = os.path.join(out, "notes.md")
    summary = write("md", md_path, {"title": "Q3 Review", "blocks": [
        {"type": "heading", "level": 2, "text": "Findings"},
        "Revenue grew 12%.",
        {"type": "bullets", "items": ["EMEA led", ["Germany +18%"]]},
        {"type": "table", "headers": ["Region", "Rev"], "rows": [["EMEA", "1.2M"]]},
    ]})
    text = open(md_path, encoding="utf-8").read()
    assert "# Q3 Review" in text and "### Findings" in text
    assert "- EMEA led" in text and "  - Germany +18%" in text
    assert "| Region | Rev |" in text and "| EMEA | 1.2M |" in text
    _raises(lambda: write_md(md_path, {"title": "x", "blocks": [{"type": "bullets", "items": []}]}), "has no items")
    # ---- Markdown -> blocks. Prose deliverables arrive this way because asking an 8B
    # model to JSON-escape a letter throws good documents away over one stray quote.
    md_src = (
        "# Late Fee Waiver\n\n"
        "Respected Sir/Madam,\n\n"
        "I can't afford the \"extra\" charges; my salary arrives late.\n\n"
        "- Account: [Account Number]\n- Amount: [Amount]\n\n"
        "| Item | Value |\n| --- | --- |\n| Due | [Date] |\n\n"
        "> Please consider this request.\n\n"
        "1. First point\n2. Second point\n\n"
        "Yours sincerely,\n[Your Name]"
    )
    spec_md = spec_from_markdown(md_src, "docx")
    assert spec_md["title"] == "Late Fee Waiver", spec_md["title"]
    kinds = [b["type"] for b in spec_md["blocks"]]
    for expected in ("paragraph", "bullets", "table", "quote", "numbered"):
        assert expected in kinds, f"{expected} missing from {kinds}"
    tbl = next(b for b in spec_md["blocks"] if b["type"] == "table")
    assert tbl["headers"] == ["Item", "Value"] and tbl["rows"] == [["Due", "[Date]"]], tbl
    # Quotes and apostrophes are exactly what broke the JSON path; they must survive.
    assert any('"extra"' in b.get("text", "") for b in spec_md["blocks"])

    md_doc = os.path.join(out, "from_md.docx")
    write_docx(md_doc, spec_md)
    reread = Document(md_doc)
    text = "\n".join(par.text for par in reread.paragraphs)
    assert "[Your Name]" in text and "salary arrives late" in text, text[:200]
    assert len(reread.tables) == 1

    # Plain prose with no Markdown at all is still a valid document, not an error.
    plain = spec_from_markdown("Just one line, no markup.", "docx")
    assert plain["blocks"][0]["text"] == "Just one line, no markup."
    assert spec_from_markdown("", "docx")["blocks"], "empty input must still yield a block"

    deck = spec_from_markdown("# Deck\n\n## One\n\n- a\n- b\n\n## Two\n\n- c", "pptx")
    assert len(deck["slides"]) == 2, deck["slides"]
    assert deck["slides"][0]["bullets"] == ["a", "b"], deck["slides"][0]
    print("PASS markdown: prose -> blocks -> docx/pptx, quotes and tables intact")

    print("PASS md:", summary)

    # ---- dispatch + atomic rewrite
    _raises(lambda: write("csv", md_path, {}), "unknown kind 'csv'", "xlsx, docx, pptx, md")
    again = write("xlsx", xlsx, {"sheets": [{"name": "Q3 Summary", "columns": ["Region"], "rows": [["EMEA"]]}]})
    wb = openpyxl.load_workbook(xlsx)
    assert wb["Q3 Summary"]["A2"].value == "EMEA", "rewrite left stale content"
    wb.close()
    assert not [f for f in os.listdir(out) if f.startswith(".ada-")], "temp file left behind"
    print("PASS rewrite:", again)

    import shutil
    shutil.rmtree(out, ignore_errors=True)


if __name__ == "__main__":
    demo()
